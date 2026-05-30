#!/usr/bin/env python3
"""Append RTOS integration slide(s) to EGH456 presentation."""

from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

PATH = "docs/EGH456_Motor_3min_Presentation.pptx"


def add_content_slide(prs, title, subtitle=""):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.clear()
    if subtitle:
        p = tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(13)
        p.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        p.space_after = Pt(10)
    return tf


def bullet(tf, text, level=0, bold=False, size=12):
    p = tf.add_paragraph()
    p.text = text
    p.level = level
    p.font.size = Pt(size)
    p.font.bold = bold
    p.space_after = Pt(3)


def main():
    prs = Presentation(PATH)

    tf = add_content_slide(
        prs,
        "RTOS Integration (2.1.6)",
        "EGH456 — FreeRTOS: ISRs, motor task, E-stop, prioritisation",
    )
    bullet(tf, "ISRs — hall speed + sensor sampling", bold=True, size=13)
    bullet(tf, "HallPortM / H / N: GPIO IRQ → g_edges + commutation (no blocking RTOS calls)", 1)
    bullet(tf, "ADC1Seq0 → xPowerRawQueue (FromISR)  ·  Timer2A → accel semaphore (FromISR)", 1)
    bullet(tf, "RPM calculated in Speed task (100 Hz), not inside hall ISR", 1)

    bullet(tf, "Motor task — 10 ms, priority idle+4", bold=True, size=13)
    bullet(tf, "State machine  ·  Ramp (500 / 1000 RPM/s)  ·  Duty from reference", 1)

    bullet(tf, "Fault / E-stop — event group xSystemEvents", bold=True, size=13)
    bullet(tf, "Sensor: EVT_ESTOP_POWER, EVT_ESTOP_ACCEL  →  Motor: ESTOP_BRAKING → FAULT", 1)
    bullet(tf, "GUI: START/STOP/ACK bits + xCommandQueue (mutex)", 1)
    bullet(tf, "Fault task (idle+1): UART logging only", 1)

    bullet(tf, "Priorities: Speed/GUI/Plot (+5) > Motor/AccSamp (+4) > Sensor (+3)", bold=True, size=13)

    tf2 = add_content_slide(prs, "RTOS Communication", "")
    p = tf2.paragraphs[0]
    p.text = (
        "Hall ISR  →  Speed task  →  Motor task (control loop)\n"
        "ADC ISR   →  queue       →  Sensor task  →  event bits  →  Motor\n"
        "Timer ISR →  semaphore   →  AccSamp task →  Sensor\n"
        "GUI       →  queue + events  →  Motor  →  queue  →  GUI"
    )
    p.font.size = Pt(13)
    p.font.name = "Courier New"
    p.alignment = PP_ALIGN.LEFT

    prs.save(PATH)
    print(f"Saved {len(prs.slides)} slides → {PATH}")


if __name__ == "__main__":
    main()
